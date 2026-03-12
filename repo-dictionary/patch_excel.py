import openpyxl
import os

BASE_DIR = r"e:\ProjectLMSGradution\LMS_Backend\repo-dictionary"
files_to_update = [
    "Vocabulary_Dictionary.xlsx",
    "Vocabulary_Full_Dictionary_NEW.xlsx",
    "Vocabulary.xlsx",
]

for filename in files_to_update:
    filepath = os.path.join(BASE_DIR, filename)
    if os.path.exists(filepath):
        print(f"Updating {filename}...")
        try:
            wb = openpyxl.load_workbook(filepath)
            if "Vocabularies" in wb.sheetnames:
                ws = wb["Vocabularies"]
            else:
                ws = wb.active

            # Check if column already exists to prevent double insertion
            header = ws.cell(row=1, column=13).value
            if header != "imageUrl":
                print(f"  Inserting imageUrl column at position 13 (M) in {filename}")
                ws.insert_cols(idx=13)

                # Copy styling from etymologyImage (col 12)
                src_cell = ws.cell(row=1, column=12)
                tgt_cell = ws.cell(row=1, column=13, value="imageUrl")

                if src_cell.has_style:
                    tgt_cell.font = openpyxl.styles.Font(
                        bold=src_cell.font.bold,
                        color=src_cell.font.color,
                        size=src_cell.font.size,
                    )
                    tgt_cell.fill = openpyxl.styles.PatternFill(
                        fill_type=src_cell.fill.fill_type,
                        start_color=src_cell.fill.start_color,
                        end_color=src_cell.fill.end_color,
                    )
                    tgt_cell.alignment = openpyxl.styles.Alignment(
                        horizontal=src_cell.alignment.horizontal,
                        vertical=src_cell.alignment.vertical,
                    )

                ws.column_dimensions["M"].width = 30
                wb.save(filepath)
                print(f"  Success: Updated {filename}")
            else:
                print(f"  Already updated: {filename} already has imageUrl at col 13")
        except Exception as e:
            print(f"  Error updating {filename}: {e}")
    else:
        print(f"File not found: {filename}")
